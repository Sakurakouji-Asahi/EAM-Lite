import io
from openpyxl import load_workbook
from django.core.exceptions import ValidationError
from .services import build_template_workbook,get_template_definition


def build_error_rows_workbook(batch):
    rows=list(batch.rows.exclude(errors_json=[]).order_by('row_number'))
    if not rows:
        raise ValidationError('没有可下载的错误行。')
    definition=get_template_definition(batch.import_type,company=batch.company,version=batch.template_version)
    if any(not row.raw_data_json for row in rows):
        raise ValidationError('包含文件结构错误，请下载原文件修正表头或格式后重新上传。')
    workbook=load_workbook(io.BytesIO(build_template_workbook(batch.import_type,company=batch.company,version=batch.template_version)))
    sheet=workbook[definition.sheet_name]
    if sheet.max_row>1: sheet.delete_rows(2,sheet.max_row-1)
    for row in rows:
        sheet.append([row.raw_data_json.get(column.name) for column in definition.columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):
                cell.data_type='s'
                cell.number_format='@'
    stream=io.BytesIO()
    workbook.save(stream)
    return stream.getvalue(),len(rows)
