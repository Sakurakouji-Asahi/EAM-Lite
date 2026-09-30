"""Report navigation, full-result summaries, exports and permission regression."""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from types import MappingProxyType

import pytest
from django.db import transaction
from django.test import RequestFactory
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.reports.excel import write_report_workbook
from apps.reports.forms import ReportFilterForm
from apps.reports.queries import ReportDataset, build_report_dataset
from apps.reports.schemas import REPORT_REGISTRY, SUPPLY_REPORT_REGISTRY, visible_report_definition
from apps.reports.summaries import summarize_report
from apps.reports.views import _dataset_context
from tests.test_sprint3_support import make_user
from tests.test_sprint4_acceptance import _base_context
from tests.test_sprint4_services import _profile_context, _confirmed_entry
from tests.test_sprint7_support import active_fixed_asset_context
from tests.test_sprint18_reports import report_context


def dataset(key, rows, *, definition=None):
    return ReportDataset(definition=definition or {**REPORT_REGISTRY, **SUPPLY_REPORT_REGISTRY}[key],
                         rows=tuple(rows), filters=MappingProxyType({}), data_snapshot_at=timezone.now())


def test_financial_summary_uses_all_rows_and_excel_matches():
    # 125 constructed rows: page three contains only the final 25 records.
    rows = [{"asset_code": str(index), "fixed_asset_category": "机器",
             "original_cost": Decimal("100.01"), "actual_accumulated_depreciation": Decimal("10.01"),
             "impairment": Decimal("0.00"), "actual_book_value": Decimal("90.00")}
            for index in range(125)]
    source = dataset("fixed_asset_detail", rows)
    context = _dataset_context(source, None, RequestFactory().get('/reports/', {'page':3}))
    assert len(context['table_rows']) == 25
    assert context['page_obj'].start_index() == 101
    totals = {card['label']:card['value'] for card in context['summary']['cards']}
    assert totals['原值合计'] == Decimal('12501.25')
    assert totals['实际累计折旧合计'] == Decimal('1251.25')
    assert totals['实际账面净值合计'] == Decimal('11250.00')
    output = BytesIO()
    write_report_workbook(source, output)
    workbook = load_workbook(output, data_only=True)
    assert workbook['固定资产明细'].max_row == 126
    headers, row = list(workbook['分类汇总'].values)
    values = dict(zip(headers, row))
    assert values['记录数'] == 125
    assert Decimal(str(values['原值'])) == totals['原值合计']


def test_quantity_units_and_hidden_cost_are_not_merged_or_exposed():
    definition = visible_report_definition('supply_stock_balance', include_supply_cost=False, include_asset_finance=False)
    source = dataset('supply_stock_balance', [
        {'warehouse_name':'仓库','unit':'件','current_quantity':Decimal('2'),'current_amount':Decimal('99')},
        {'warehouse_name':'仓库','unit':'kg','current_quantity':Decimal('3'),'current_amount':Decimal('88')},
    ], definition=definition)
    summary = summarize_report(source)
    assert {row['unit']:row['current_quantity'] for row in summary['rows']} == {'件':Decimal('2'),'kg':Decimal('3')}
    assert all(column.key != 'current_amount' for column in summary['columns'])
    assert summary['cards'] == [{'label':'记录数','value':2,'kind':'integer'}]
    output = BytesIO()
    write_report_workbook(source, output)
    workbook = load_workbook(output, data_only=True)
    assert not any('金额' in str(cell) for row in workbook['分类汇总'].values for cell in row)


def test_summary_retains_missing_values_and_zero_and_signed_reversal():
    summary = summarize_report(dataset('fixed_asset_detail', [
        {'fixed_asset_category':'机器','original_cost':Decimal('0.00')},
        {'fixed_asset_category':'机器','original_cost':None},
    ]))
    assert summary['rows'][0]['original_cost'] == Decimal('0.00')
    assert summary['rows'][0]['actual_book_value'] is None
    assert '原值有 1 条缺失' in summary['missing'][0]
    movements = summarize_report(dataset('supply_stock_ledger', [
        {'movement_type':'领用','unit':'件','quantity_delta':Decimal('-3'),'amount_delta':Decimal('-30'),'amount_after':Decimal('70')},
        {'movement_type':'冲销','unit':'件','quantity_delta':Decimal('3'),'amount_delta':Decimal('30'),'amount_after':Decimal('100')},
    ]))
    assert movements['cards'][1]['value'] == Decimal('0')
    assert 'amount_after' not in {column.key for column in movements['columns']}


def test_repeated_net_issue_balances_and_management_subtotals_are_not_summed():
    summary = summarize_report(dataset('supply_issue_detail', [
        {'business_type':'领用','unit':'件','quantity':Decimal('5'),'amount':Decimal('50'),'current_net_amount':Decimal('30')},
        {'business_type':'退回','unit':'件','quantity':Decimal('2'),'amount':Decimal('20'),'current_net_amount':Decimal('30')},
    ]))
    assert len(summary['rows']) == 2
    assert len(summary['cards']) == 1
    assert 'current_net_amount' not in {column.key for column in summary['columns']}
    management = summarize_report(dataset('supply_management_amount', [
        {'component':'库存','supply_amount':Decimal('40')},
        {'component':'保管','supply_amount':Decimal('60')},
        {'component':'小计','supply_amount':Decimal('100')},
    ]))
    assert len(management['cards']) == 1
    assert management['columns'] == ()


def test_superseded_depreciation_plans_are_summarized_separately():
    summary = summarize_report(dataset('depreciation_schedule', [
        {'schedule_status':'计划','period_start':date(2026,9,1),'period_end':date(2026,10,1),'theoretical_amount':Decimal('120.00')},
        {'schedule_status':'已替代','period_start':date(2026,9,1),'period_end':date(2026,10,1),'theoretical_amount':Decimal('100.00')},
    ]))
    assert len(summary['cards']) == 1
    assert {row['schedule_status']:row['theoretical_amount'] for row in summary['rows']} == {'计划':Decimal('120.00'),'已替代':Decimal('100.00')}


@pytest.mark.django_db(transaction=True)
def test_report_opens_with_data_searches_equipment_and_reset_keeps_report(client):
    context, asset, *_ = active_fixed_asset_context('RPTSEARCH')
    asset.equipment_number = 'SB-2026-009'
    asset.save(update_fields=['equipment_number'])
    client.force_login(context['finance'])
    home = client.get(reverse('reports:report-center'))
    assert home.status_code == 200 and home.context['dataset'].row_count == 1
    assert 'period_start' not in home.context['form'].fields
    assert 'maintenance_due_scope' not in home.context['form'].fields
    search = client.get(reverse('reports:report-center'), {'report_type':'fixed_asset_detail','q':'sb-2026'})
    assert search.status_code == 200 and search.context['dataset'].row_count == 1
    assert search.context['dataset'].rows[0]['equipment_number'] == 'SB-2026-009'
    assert search.context['reset_url'].endswith('report_type=fixed_asset_detail')
    assert client.get(reverse('reports:report-center'), {'report_type':'fixed_asset_detail','q':'不存在'}).context['dataset'].row_count == 0
    assert client.get(reverse('reports:report-center'), {'page_size':'1000'}).status_code == 400
    invalid = client.get(reverse('reports:report-center'), {'report_type':'asset_ledger','period_start':'2026-09-01','period_end':'2026-09-30'})
    assert invalid.status_code == 400
    assert '不适用于当前报表' in invalid.content.decode()


@pytest.mark.django_db(transaction=True)
def test_report_export_is_all_matching_rows_and_cost_access_is_unchanged(client, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    context, asset, *_ = active_fixed_asset_context('RPTEXPORT')
    client.force_login(context['finance'])
    response = client.post(reverse('reports:report-export'), {'report_type':'fixed_asset_detail','q':asset.asset_code,'idempotency_key':'report-workspace-export'})
    assert response.status_code == 302
    from apps.reports.models import ExportLog
    from django.core.files.storage import default_storage
    export = ExportLog.objects.get(idempotency_key='report-workspace-export')
    assert export.row_count == 1
    detail = client.get(response.headers['Location'])
    assert detail.status_code == 200
    assert detail.context['display_filters']['关键词'] == asset.asset_code
    assert 'q=' in detail.context['return_report_url']
    assert '返回当前报表' in detail.content.decode()
    with default_storage.open(export.output_attachment.storage_key, 'rb') as file:
        workbook = load_workbook(file, data_only=True)
    assert '分类汇总' in workbook.sheetnames
    client.force_login(context['equipment'])
    forbidden = client.get(reverse('reports:report-center'), {'report_type':'fixed_asset_detail'})
    assert forbidden.status_code == 403
    allowed = client.get(reverse('reports:report-center'))
    assert '原值' not in allowed.content.decode()
    assert 'no-store' in allowed.headers['Cache-Control']


@pytest.mark.django_db(transaction=True)
def test_monthly_components_separate_opening_from_posted_depreciation():
    from apps.finance.models import DepreciationEntry
    from apps.finance.services import create_value_adjustment, reverse_value_adjustment
    company, actor, _management, _admin, asset, _finance, profile = _profile_context()
    start = timezone.localdate().replace(day=1)
    DepreciationEntry.objects.create(
        company=company, asset=asset, depreciation_profile=profile, entry_date=start,
        period_start=start, period_end=start + timedelta(days=1), source_type='opening', opening_profile=profile,
        amount=Decimal('1000.00'), accumulated_depreciation_after=Decimal('1000.00'),
        book_value_after=Decimal('11000.00'), posted_by=actor, posted_at=timezone.now())
    adjustment = create_value_adjustment(actor=actor, asset=asset, adjustment_type='depreciation_adjustment',
                                        amount=Decimal('5.12'), effective_date=start, reason='报表构造样例')
    reversal = reverse_value_adjustment(actor=actor, adjustment=adjustment, reason='核对冲销列')
    current_profile = asset.depreciation_profiles.get(status='active')
    with transaction.atomic():
        _confirmed_entry(profile=current_profile, start=start, amount=Decimal('19.73'), actor=actor)
    result = build_report_dataset(actor=actor, company=company, report_key='monthly_depreciation', filters={'include_drafts':True})
    summary = summarize_report(result)
    row = next(row for row in summary['rows'] if row['period_start'] == start)
    assert row['opening_amount'] == Decimal('1000.00')
    assert row['depreciation_amount'] == Decimal('19.73')
    assert row['adjustment_amount'] == Decimal('5.12')
    reversal_row = next(item for item in summary['rows'] if item['period_start'] == reversal.effective_date.replace(day=1))
    assert reversal_row['reversal_amount'] == Decimal('-5.12')
    assert row['actual_amount'] == Decimal('1024.85') + row['reversal_amount']
    assert sum(item['actual_amount'] for item in summary['rows']) == Decimal('1019.73')


@pytest.mark.django_db(transaction=True)
def test_supply_keyword_default_period_and_summary_match_source(client):
    context = report_context()
    client.force_login(context['warehouse_user'])
    response = client.get(reverse('reports:supply-report-detail', args=['supply_stock_balance']), {'q':'PAPER'})
    assert response.status_code == 200
    assert response.context['dataset'].row_count == 1
    row = response.context['summary']['rows'][0]
    assert row['current_quantity'] == Decimal('7.0000')
    assert row['current_amount'] == Decimal('700.00')
    movement = client.get(reverse('reports:supply-report-detail', args=['supply_stock_movement']))
    assert movement.status_code == 200 and movement.context['dataset'] is not None
    assert movement.context['form'].cleaned_data['date_from'] == timezone.localdate().replace(day=1)


def test_export_summary_reads_stream_only_once_and_separates_same_named_departments():
    class OnceOnlyRows:
        def __len__(self):
            return 2

        def __iter__(self):
            assert not getattr(self, 'consumed', False)
            self.consumed = True
            yield {'asset_code':'A','department':'生产部','_summary_identity':{'department':1}}
            yield {'asset_code':'B','department':'生产部','_summary_identity':{'department':2}}

    source = ReportDataset(definition=REPORT_REGISTRY['department_assets'], rows=OnceOnlyRows(),
                           filters={}, data_snapshot_at=timezone.now())
    output = BytesIO()
    write_report_workbook(source, output)
    workbook = load_workbook(output, data_only=True)
    assert workbook['分类汇总'].max_row == 3
    assert workbook['部门资产'].max_row == 3


@pytest.mark.django_db(transaction=True)
def test_tplus_assets_and_entries_have_independent_full_pagination(client, monkeypatch):
    from apps.reports.queries import TplusDataset
    from apps.reports.schemas import TPLUS_TOTAL_METRICS
    context = _base_context('RPTPAGES')
    client.force_login(context['finance'])
    source = TplusDataset(
        definition=REPORT_REGISTRY['tplus_reconciliation'],
        asset_rows=tuple({'asset_code':str(index)} for index in range(130)),
        entry_rows=tuple({'asset_code':str(index)} for index in range(155)),
        filters={}, data_snapshot_at=timezone.now(),
        totals={key:Decimal('123.45') for key in TPLUS_TOTAL_METRICS})
    monkeypatch.setattr('apps.reports.views.build_tplus_dataset', lambda **kwargs: source)
    result = client.get(reverse('reports:tplus-export'), {'period':'2026-09','asset_page':3,'entry_page':4})
    assert result.status_code == 200
    assert len(result.context['asset_preview_rows']) == 30
    assert len(result.context['entry_preview_rows']) == 5
    assert result.context['asset_page_obj'].start_index() == 101
    assert result.context['entry_page_obj'].start_index() == 151
    assert all(value == Decimal('123.45') for _, value in result.context['total_rows'])
