from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.masterdata.models import UserDepartmentScope
from apps.supplies.models import SupplyStockLedger
from apps.supplies.services import (
    post_supply_document,
    publish_supply_count_task,
    record_supply_count,
)
from tests.test_sprint15_services import supply_context
from tests.test_sprint15_support import (
    make_department,
    make_employee,
    make_issue_document,
    make_user,
    seed_supply_stock,
)
from tests.test_sprint17_services import make_count


pytestmark = pytest.mark.django_db(transaction=True)


@pytest.fixture
def count_role_scope():
    company, warehouse_actor, department_b, mine, warehouse, _, _, item = supply_context()
    actor = make_user("count-scope-mixed", "employee", "department_manager")
    mine.user = actor
    mine.save(update_fields=["user", "updated_at"])
    other = make_employee(company, department_b, "B-OTHER")
    department_a = make_department(company, "MANAGED-A")
    managed_employees = [
        make_employee(company, department_a, "A-FIRST"),
        make_employee(company, department_a, "A-SECOND"),
    ]
    UserDepartmentScope.objects.create(
        company=company,
        user=actor,
        department=department_a,
        include_descendants=False,
        assigned_by=warehouse_actor,
    )
    seed_supply_stock(
        actor=warehouse_actor,
        company=company,
        warehouse=warehouse,
        item=item,
        quantity="12",
        unit_cost="80",
        key="count-role-scope-opening",
    )
    for employee, quantity in zip(
        [mine, other, *managed_employees], ["1", "2", "3", "4"], strict=True
    ):
        document = make_issue_document(
            actor=warehouse_actor,
            company=company,
            warehouse=warehouse,
            item=item,
            department=employee.department,
            employee=employee,
            quantity=quantity,
            key=f"count-role-scope-issue-{employee.employee_no}",
        )
        post_supply_document(document=document, actor=warehouse_actor)

    equipment_actor = make_user("count-scope-equipment", "equipment")
    tasks = []
    for department, key in ((department_a, "managed-a"), (department_b, "personal-b")):
        task = make_count(
            actor=equipment_actor,
            company=company,
            domain="custody",
            department=department,
            key=f"count-role-scope-{key}",
        )
        publish_supply_count_task(task=task, actor=equipment_actor)
        tasks.append(task)
    task_a, task_b = tasks
    own_line = task_b.lines.get(custody__employee=mine)
    other_line = task_b.lines.get(custody__employee=other)
    record_supply_count(
        actor=equipment_actor,
        line=other_line,
        counted_quantity=Decimal("1"),
        remark="其他责任人的盘点差异",
    )
    return SimpleNamespace(
        actor=actor,
        task_a=task_a,
        task_b=task_b,
        own_line=own_line,
        other_line=other_line,
        mine=mine,
        other=other,
        managed_employees=managed_employees,
    )


def _page_line_ids(response):
    return {row.pk for row in response.context["page_obj"]}


def _workbook_rows(content):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        return list(workbook["盘点表"].iter_rows(values_only=True))[1:]
    finally:
        workbook.close()


def test_mixed_roles_show_only_personal_rows_and_summary_outside_managed_department(
    client, count_role_scope
):
    scope = count_role_scope
    client.force_login(scope.actor)
    response = client.get(reverse("supplies:count-task-detail", args=[scope.task_b.pk]))
    assert response.status_code == 200
    assert _page_line_ids(response) == {scope.own_line.pk}
    summary = response.context["count_summary"]
    assert {key: summary[key] for key in ("total", "recorded", "unrecorded", "different")} == {
        "total": 1,
        "recorded": 0,
        "unrecorded": 1,
        "different": 0,
    }
    assert response.context["page_obj"].paginator.count == 1
    assert not response.context["show_cost"]
    html = response.content.decode()
    assert scope.mine.name in html
    assert scope.other.name not in html
    assert "其他责任人的盘点差异" not in html
    assert "应盘金额" not in html
    assert "80.00" not in html and "160.00" not in html


@pytest.mark.parametrize(
    "query",
    [
        {"q": "B-OTHER", "page_size": "25", "page": "2"},
        {"row_view": "different", "page_size": "25", "page": "2"},
    ],
)
def test_personal_scope_applies_before_filtering_pagination_and_summary(
    client, count_role_scope, query
):
    scope = count_role_scope
    client.force_login(scope.actor)
    response = client.get(
        reverse("supplies:count-task-detail", args=[scope.task_b.pk]), query
    )
    assert response.status_code == 200
    assert response.context["page_obj"].paginator.count == 0
    assert _page_line_ids(response) == set()
    assert response.context["count_summary"]["total"] == 1
    assert response.context["count_summary"]["different"] == 0
    assert scope.other.name not in response.content.decode()


def test_personal_bulk_entry_and_export_do_not_include_other_employees(
    client, count_role_scope
):
    scope = count_role_scope
    client.force_login(scope.actor)
    bulk = client.get(reverse("supplies:count-task-bulk-entry", args=[scope.task_b.pk]))
    assert bulk.status_code == 200
    assert {form.line.pk for form in bulk.context["row_forms"]} == {scope.own_line.pk}
    assert bulk.context["count_summary"]["total"] == 1
    assert all("adjustment_unit_cost" not in form.fields for form in bulk.context["row_forms"])

    exported = client.get(
        reverse("supplies:count-sheet", args=[scope.task_b.pk]), {"download": "1"}
    )
    assert exported.status_code == 200
    rows = _workbook_rows(exported.content)
    assert len(rows) == 1
    assert rows[0][10] == scope.mine.name
    assert rows[0][4] == "1.0000"
    assert rows[0][7:9] == (None, None)


def test_personal_recording_is_idempotent_and_other_line_is_denied(client, count_role_scope):
    scope = count_role_scope
    client.force_login(scope.actor)
    own_url = reverse("supplies:count-line-record", args=[scope.task_b.pk, scope.own_line.pk])
    other_url = reverse("supplies:count-line-record", args=[scope.task_b.pk, scope.other_line.pk])
    assert client.get(own_url).status_code == 200
    assert client.get(other_url).status_code == 403
    assert client.post(other_url, {"counted_quantity": "0", "remark": "越权修改"}).status_code == 403
    audit_before = AuditLog.objects.filter(action="supply_count_record").count()
    ledger_before = SupplyStockLedger.objects.count()
    data = {"counted_quantity": "0", "remark": "本人实盘无实物", "expected_counted_at": ""}
    assert client.post(own_url, data).status_code == 302
    scope.own_line.refresh_from_db()
    saved_at = scope.own_line.counted_at
    assert scope.own_line.counted_quantity == Decimal("0")
    assert scope.own_line.counted_by_id == scope.actor.pk
    assert AuditLog.objects.filter(action="supply_count_record").count() == audit_before + 1
    assert client.post(own_url, data).status_code == 302
    scope.own_line.refresh_from_db()
    scope.other_line.refresh_from_db()
    assert scope.own_line.counted_at == saved_at
    assert scope.other_line.counted_quantity == Decimal("1")
    assert scope.other_line.remark == "其他责任人的盘点差异"
    assert AuditLog.objects.filter(action="supply_count_record").count() == audit_before + 1
    assert SupplyStockLedger.objects.count() == ledger_before


def test_managed_department_keeps_all_rows_visible_and_recordable(client, count_role_scope):
    scope = count_role_scope
    client.force_login(scope.actor)
    expected_ids = set(scope.task_a.lines.values_list("pk", flat=True))
    detail = client.get(reverse("supplies:count-task-detail", args=[scope.task_a.pk]))
    assert detail.status_code == 200
    assert _page_line_ids(detail) == expected_ids
    assert detail.context["count_summary"]["total"] == 2
    assert all(row["can_record"] for row in detail.context["line_rows"])
    url = reverse("supplies:count-task-bulk-entry", args=[scope.task_a.pk])
    bulk = client.get(url)
    assert bulk.status_code == 200
    data = {"line_manifest": bulk.context["line_manifest"]}
    for form in bulk.context["row_forms"]:
        data[f"line-{form.line.pk}-counted_quantity"] = str(form.line.expected_quantity)
        data[f"line-{form.line.pk}-expected_counted_at"] = ""
    assert client.post(url, data).status_code == 302
    for line in scope.task_a.lines.all():
        assert line.counted_quantity == line.expected_quantity
        assert line.counted_by_id == scope.actor.pk


def test_removing_employee_role_revokes_task_and_previously_exported_sheet(
    client, count_role_scope
):
    scope = count_role_scope
    client.force_login(scope.actor)
    sheet_url = reverse("supplies:count-sheet", args=[scope.task_b.pk])
    exported = client.get(sheet_url, {"download": "1"})
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content))
    workbook["盘点表"]["F2"] = "0"
    workbook["盘点表"]["G2"] = "权限撤销后回传旧校验码"
    completed_sheet = BytesIO()
    workbook.save(completed_sheet)
    workbook.close()
    scope.actor.groups.remove(Group.objects.get(name="employee"))
    assert client.get(reverse("supplies:count-task-detail", args=[scope.task_b.pk])).status_code == 404
    assert client.get(reverse("supplies:count-task-bulk-entry", args=[scope.task_b.pk])).status_code == 404
    assert client.get(sheet_url, {"download": "1"}).status_code == 404
    audit_before = AuditLog.objects.filter(action="supply_count_record").count()
    upload = SimpleUploadedFile(
        "previously-authorized-count.xlsx",
        completed_sheet.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert client.post(sheet_url, {"file": upload}).status_code == 404
    scope.own_line.refresh_from_db()
    assert scope.own_line.counted_quantity is None
    assert AuditLog.objects.filter(action="supply_count_record").count() == audit_before
    assert client.get(reverse("supplies:count-task-detail", args=[scope.task_a.pk])).status_code == 200


@pytest.mark.parametrize("role", ["warehouse", "equipment", "management"])
def test_full_view_roles_keep_all_department_rows_and_cost_visibility(
    client, count_role_scope, role
):
    scope = count_role_scope
    actor = make_user(f"count-full-view-{role}", role)
    client.force_login(actor)
    response = client.get(reverse("supplies:count-task-detail", args=[scope.task_b.pk]))
    assert response.status_code == 200
    assert _page_line_ids(response) == {scope.own_line.pk, scope.other_line.pk}
    assert response.context["count_summary"]["total"] == 2
    assert response.context["count_summary"]["different"] == 1
    assert response.context["show_cost"]
    assert scope.other.name in response.content.decode()
    assert "应盘金额" in response.content.decode()
