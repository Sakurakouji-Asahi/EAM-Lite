"""Constructed business cases for monthly entry and searchable workspaces."""
import io
from datetime import date,timedelta
from decimal import Decimal
from types import SimpleNamespace

import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.finance.models import AssetWorkUsage,DepreciationEntry,TheoreticalDepreciationRun
from apps.finance.usage_workspace import save_usage_page,monthly_profiles
from apps.finance.services import generate_depreciation_batch,confirm_depreciation_batch,run_theoretical_depreciation
from apps.reports.models import ReportPreset
from apps.reports.relative_periods import apply_relative_period
from tests.test_correction_finance_services import _custom_profile_context
from tests.test_sprint7_support import active_asset_context
from tests.test_sprint9_support import maintenance_context
from tests.test_sprint9_services import _complete
from tests.test_sprint3_support import make_user

pytestmark=pytest.mark.django_db(transaction=True)


def test_monthly_usage_keeps_blank_zero_and_retries_and_correctly_uses_period_version(client):
    company,actor,management,admin,asset,finance,profile=_custom_profile_context(method='units_of_production',start_date=date(2026,8,1))
    client.force_login(actor)
    url=reverse('finance:monthly-usage')+'?month=2026-09&state=missing'
    page=client.get(url)
    assert page.status_code==200 and page.context['missing']==1
    data={'manifest':page.context['manifest'],f'usage-{profile.pk}-units':'',f'usage-{profile.pk}-remark':''}
    assert client.post(url,data).status_code==302
    assert not AssetWorkUsage.objects.exists()
    data[f'usage-{profile.pk}-units']='0'
    assert client.post(url,data).status_code==302
    assert client.post(url,data).status_code==302
    assert AssetWorkUsage.objects.get().current_units==Decimal('0')
    assert client.get(url).context['missing']==0
    data[f'usage-{profile.pk}-units']='1'
    rejected=client.post(url,data)
    assert rejected.status_code==200 and '已提交其他内容' in rejected.content.decode()
    assert AssetWorkUsage.objects.count()==1
    client.force_login(management)
    assert client.get(url).status_code==403


def test_usage_cap_is_disclosed_and_bulk_failure_does_not_write():
    company,actor,_,_,asset,_,profile=_custom_profile_context(method='units_of_production',start_date=date(2026,8,1))
    entries=[{'profile':str(profile.pk),'units':Decimal('120'),'remark':'实际运行'}]
    assert save_usage_page(actor=actor,company=company,start=date(2026,9,1),end=date(2026,10,1),entries=entries,key='cap-usage')==(1,1)
    usage=AssetWorkUsage.objects.get()
    assert usage.current_units==usage.closing_accumulated_units==Decimal('100')
    assert '120' in usage.remark and '封顶' in usage.remark
    assert save_usage_page(actor=actor,company=company,start=date(2026,9,1),end=date(2026,10,1),entries=entries,key='cap-usage')==(0,1)
    with pytest.raises(ValidationError):
        save_usage_page(actor=actor,company=company,start=date(2026,10,1),end=date(2026,11,1),entries=[*entries,{'profile':'00000000-0000-0000-0000-000000000001','units':Decimal('1'),'remark':''}],key='outside-usage')
    assert AssetWorkUsage.objects.count()==1


def test_plan_filters_validate_scope_and_history_pages(client):
    from apps.maintenance.services import update_maintenance_plan
    ctx=maintenance_context('WORK9MAINT')
    plan=ctx['plan'];today=timezone.localdate()
    ctx['plan']=update_maintenance_plan(actor=ctx['equipment'],plan=plan,name=plan.name,cycle_value=1,cycle_unit='day',
        responsible_employee=ctx['responsible'],advance_notice_days=3,standard_content=plan.standard_content,first_due_date=today-timedelta(days=40))
    for index in range(27):
        ctx['plan'].refresh_from_db()
        _complete(ctx,f'work9-complete-{index}',completed_date=today-timedelta(days=40-index),remark=f'核对记录 {index}')
    client.force_login(ctx['equipment'])
    page=client.get(reverse('maintenance:plan-detail',args=[plan.pk]),{'page':2})
    assert page.status_code==200 and page.context['page_obj'].paginator.count==27 and len(page.context['records'])==2
    assert client.get(reverse('maintenance:plan-list'),{'status':'invented'}).status_code==400
    due=client.get(reverse('maintenance:due-list'),{'responsible_employee':ctx['responsible'].pk,'due_scope':'overdue'})
    assert due.status_code==200 and due.context['page_obj'].paginator.count==1
    assert client.get(reverse('maintenance:due-list'),{'department':'99999999'}).status_code==400
    client.force_login(make_user('work9-maint-outside','employee'))
    assert client.get(reverse('maintenance:due-list'),{'responsible_employee':ctx['responsible'].pk}).status_code==400


def test_finance_history_filters_do_not_change_whole_asset_balance(client):
    company,actor,_,_,asset,_,profile=_custom_profile_context(method='straight_line',start_date=date(2024,1,1))
    for index in range(26):
        start=date(2024+index//12,index%12+1,1)
        end=date(start.year+1,1,1) if start.month==12 else date(start.year,start.month+1,1)
        batch=generate_depreciation_batch(actor=actor,company=company,period_start=start,period_end=end,idempotency_key=f'history9-{index}')
        confirm_depreciation_batch(actor=actor,batch=batch,reason="构造历史核对")
    client.force_login(actor)
    url=reverse('finance:asset-finance-detail',args=[asset.pk])
    page=client.get(url,{'entry_page':2})
    assert page.context['entry_page'].paginator.count==26 and len(page.context['entries'])==1
    narrowed=client.get(url,{'entry-year':2024,'entry-source':'batch'})
    assert narrowed.context['entry_page'].paginator.count==12
    assert page.context['actual_ad']==narrowed.context['actual_ad']==Decimal('4940.00')
    assert narrowed.context['book_value']==Decimal('7060.00')


def test_error_row_workbook_is_standard_and_preserves_bad_values(client,settings,tmp_path):
    from tests.test_sprint1_imports import setup_data,workbook_file
    from apps.imports.services import upload_and_validate_import
    settings.MEDIA_ROOT=tmp_path
    company,actor=setup_data('system_admin')
    batch=upload_and_validate_import(actor=actor,company=company,import_type='department',
        uploaded_file=workbook_file('department',[['GOOD','正常','','','是'],['BAD','未匹配','MISSING','','是']]),idempotency_key='errors9')
    client.force_login(actor)
    response=client.get(reverse('imports:error-rows',args=[batch.pk]))
    assert response.status_code==200
    workbook=load_workbook(io.BytesIO(response.content),data_only=False)
    sheet=workbook['部门导入']
    assert sheet.max_row==2 and sheet.cell(2,1).value=='BAD' and sheet.cell(2,3).value=='MISSING'
    sheet.cell(2,3).value='';stream=io.BytesIO();workbook.save(stream)
    corrected=upload_and_validate_import(actor=actor,company=company,import_type='department',
        uploaded_file=SimpleUploadedFile('corrected.xlsx',stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),idempotency_key='corrected9')
    assert corrected.rows.get().errors_json==[] and batch.rows.count()==2


def test_dynamic_query_dates_cross_year_without_saving_old_data(client,monkeypatch):
    monkeypatch.setattr('apps.reports.relative_periods.timezone',SimpleNamespace(localdate=lambda:date(2027,1,15)))
    assert apply_relative_period('monthly_depreciation',{'report_type':'monthly_depreciation','period_start':'2000-01-01','period_end':'2000-01-31','_relative_period':'previous_month'})=={
        'report_type':'monthly_depreciation','period_start':'2026-12-01','period_end':'2026-12-31'}
    ctx,asset,_=active_asset_context('RELATIVE9')
    client.force_login(ctx['finance'])
    page=client.get(reverse('reports:report-center'),{'report_type':'monthly_depreciation','period_start':'2026-09-01','period_end':'2026-09-30'})
    result=client.post(reverse('reports:preset-save'),{'preset_token':page.context['preset_token'],'name':'上月折旧','relative_period':'previous_month'})
    assert result.status_code==302 and '2026-12-01' in result.url
    preset=ReportPreset.objects.get()
    assert preset.query['period_start']=='2026-09-01' and preset.query['_relative_period']=='previous_month'


def test_asset_export_preview_second_page_keeps_full_export_filters(client,monkeypatch):
    ctx,asset,_=active_asset_context('EXPORT9')
    rows=[{'asset_code':f'ROW{index:03d}','asset_name':'分页样例'} for index in range(105)]
    monkeypatch.setattr('apps.reports.queries.build_report_dataset',lambda **kwargs:SimpleNamespace(rows=rows,row_count=105,data_snapshot_at=timezone.now()))
    client.force_login(ctx['finance'])
    page=client.get(reverse('assets:asset-list-export'),{'page':2})
    assert page.status_code==200 and len(page.context['preview_rows'])==5
    assert page.context['dataset'].row_count==105 and 'page' not in page.context['filters']


def test_master_tree_paths_and_deactivation_references_respect_current_scope(client):
    ctx,asset,_=active_asset_context('REFERENCES9')
    client.force_login(ctx['admin'])
    page=client.get(reverse('masterdata:location-list'),{'q':asset.location.name})
    assert page.status_code==200
    row=next(row for row in page.context['rows'] if row['object'].pk==asset.location_id)
    assert asset.location.parent.name in row['ancestor_path']
    detail=client.get(reverse('masterdata:location-detail',args=[asset.location_id]))
    assert detail.status_code==200 and asset.asset_code in detail.content.decode()


def test_audit_search_resolves_business_code_without_bypassing_scope(client):
    ctx,asset,_=active_asset_context('AUDITSEARCH9')
    client.force_login(ctx['finance'])
    url=reverse('audit:log-list')
    page=client.get(url,{'q':asset.asset_code})
    assert page.status_code==200 and page.context['page_obj'].paginator.count>0
    labels=client.get(url,{'q':'创建资产草稿'})
    assert labels.status_code==200
    client.force_login(make_user('audit-search-employee','employee'))
    assert client.get(url,{'q':asset.asset_code}).status_code==403


def test_theory_history_compares_saved_parameters_without_posting(client):
    company,actor,_,_,asset,finance,profile=_custom_profile_context(method='straight_line',start_date=date(2026,9,1))
    client.force_login(actor)
    for cost in ('12000','14000'):
        result=client.post(reverse('finance:theoretical-run',args=[asset.pk]),{'as_of_date':'2026-09-30','idempotency_key':'theory9-'+cost,
            'original_cost':cost,'method':'straight_line','posting_period':'monthly','commissioning_date':'2026-09-01','start_rule':'specified_date','specified_start':'2026-09-01','useful_life_months':'60','salvage_mode':'rate','salvage_percent':'5',
            'opening_actual_accumulated_depreciation':'0','opening_impairment':'0','opening_book_value':cost,'actual_continuation_date':'2026-09-01'})
        assert result.status_code==302
    runs=list(TheoreticalDepreciationRun.objects.filter(asset=asset))
    url=reverse('finance:theoretical-history',args=[asset.pk])
    page=client.get(url,{'compare_a':runs[0].pk,'compare_b':runs[1].pk})
    assert page.status_code==200 and page.context['page_obj'].paginator.count==2
    assert any(row['changed'] for row in page.context['comparison'])
    assert not DepreciationEntry.objects.exists()
    finance.refresh_from_db();assert finance.original_cost==Decimal('12000')
    assert client.get(url,{'compare_a':'00000000-0000-0000-0000-000000000001'}).status_code==400
