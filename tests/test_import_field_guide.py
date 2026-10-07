"""Downloaded template help and upload-page help share the real column contract."""

from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.imports.field_guide import template_field_guide
from apps.imports.services import (
    ASSET_FINANCE_KEYS, TEMPLATE_REGISTRY, _load_rows,
    build_template_workbook, get_template_definition,
)
from tests.test_sprint1_imports import setup_data
from tests.test_sprint3_support import make_category, make_custom_field


@pytest.mark.parametrize("import_type", TEMPLATE_REGISTRY)
def test_template_guide_preserves_headers_version_and_parseability(import_type):
    definition = get_template_definition(import_type)
    content = build_template_workbook(import_type)
    rows, errors = _load_rows(content, definition)
    assert rows == [] and errors == []
    workbook = load_workbook(BytesIO(content))
    instructions = workbook["填写说明"]
    assert instructions["A2"].value == "模板版本"
    assert instructions["B2"].value == definition.version
    assert tuple(cell.value for cell in workbook[definition.sheet_name][1]) == definition.headers
    guide_start = next(row[0].row for row in instructions if row[0].value == "字段填写指引")
    actual = [(row[0].value, row[1].value) for row in instructions.iter_rows(min_row=guide_start + 1)]
    guide = template_field_guide(definition, finance_keys=ASSET_FINANCE_KEYS if import_type == "asset_initialization" else ())
    assert actual == [(field["name"], f"{'每行必填' if field['required'] else '按业务填写'}；{field['hint']}") for field in guide]
    assert instructions.cell(guide_start + 1, 2).alignment.wrap_text
    workbook.close()


def test_guidance_distinguishes_single_asset_and_quantity_imports():
    asset = {field["key"]: field for field in template_field_guide(get_template_definition("asset_initialization"), finance_keys=ASSET_FINANCE_KEYS)}
    stock = {field["key"]: field for field in template_field_guide(get_template_definition("opening_stock"))}
    assert "精确为 1" in asset["quantity"]["hint"]
    assert "大于 0" in stock["quantity"]["hint"] and "4 位小数" in stock["quantity"]["hint"]
    assert "6 位小数" in stock["unit_cost"]["hint"]
    assert "为 0 时必填" in stock["zero_cost_reason"]["hint"]
    assert "没有财务权限时留空" in asset["original_cost"]["hint"]
    assert "财务权限" not in asset["asset_name"]["hint"]


@pytest.mark.django_db
def test_upload_field_guide_includes_active_custom_columns_and_current_company(client):
    company, actor = setup_data("equipment")
    category = make_category(company, "GUIDE-CATEGORY")
    active = make_custom_field(company, category, "GUIDE-ACTIVE", "text")
    make_custom_field(company, category, "GUIDE-DISABLED", "text", active=False)
    client.force_login(actor)
    response = client.get(reverse("imports:upload", args=["asset_initialization"]))
    assert response.status_code == 200
    expected = get_template_definition("asset_initialization", company=company)
    assert [field["name"] for field in response.context["field_guide"]] == list(expected.headers)
    assert f"自定义:{active.code}" in response.content.decode()
    assert "自定义:GUIDE-DISABLED" not in response.content.decode()
    assert response.context["max_import_rows"] == 10000
    assert company.short_name in response.content.decode()
    assert "data-import-field-search" in response.content.decode()
    assert "data-import-required-only" in response.content.decode()
